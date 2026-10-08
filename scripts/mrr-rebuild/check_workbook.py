#!/usr/bin/env python3
"""Data checks on a workbook written by consolidate_stripe.py.

Usage: check_workbook.py OUTPUT.xlsx STRIPE.csv
Exits non-zero if any check fails.
"""
import sys
from collections import Counter, defaultdict

import openpyxl

import consolidate_stripe as cs

TOL = 0.05


def grid(ws, first_col=8):
    """{(row, col): value} for fee rows only."""
    out = {}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[1] is not None:  # facility header row
            continue
        for c, v in enumerate(row[first_col - 1:], start=first_col):
            if isinstance(v, (int, float)) and v:
                out[(r, c)] = v
    return out


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def main(path, stripe_csv):
    wb = openpyxl.load_workbook(path, read_only=True)
    results = []

    def check(name, ok, detail=""):
        results.append(ok)
        print(f"{'PASS' if ok else 'FAIL'}  {name}{'  -- ' + detail if detail else ''}")

    g = {t: grid(wb[t]) for t in ("HubSpot", "Stripe", "Consolidated", "Synced Duplicates")}
    keys = set().union(*g.values())
    bad = [k for k in keys
           if abs(g["HubSpot"].get(k, 0) - g["Synced Duplicates"].get(k, 0) + g["Stripe"].get(k, 0)
                  - g["Consolidated"].get(k, 0)) > TOL]
    check("Every cell: Consolidated = HubSpot - Synced Duplicates + Stripe", not bad,
          f"{len(bad)} cells off, e.g. {bad[:3]}" if bad else f"{len(keys)} cells")

    hs_detail, st_detail = table(wb["HubSpot Detail"]), table(wb["Stripe Detail"])
    for title, detail, amount in (("HubSpot", hs_detail, "Amount"), ("Stripe", st_detail, "Amount")):
        diff = sum(g[title].values()) - sum(r[amount] for r in detail)
        check(f"{title} grid total = {title} detail total", abs(diff) < 1, f"diff {diff:.2f}")

    lines = cs.read_stripe(stripe_csv)
    billable = [r for r in lines if r["invoice_status"] not in cs.SKIP_STRIPE_STATUSES]
    not_placed = table(wb["Stripe Not Placed"])
    placed_ids = Counter(r["Line Item ID"] for r in st_detail)
    np_total = sum(r["Amount"] for r in not_placed)
    diff = sum(r["net"] for r in billable) - sum(r["Amount"] for r in st_detail) - np_total
    check("Stripe billable total = placed + not placed", abs(diff) < 1, f"diff {diff:.2f}")
    billable_ids = {r["line_item_id"] for r in billable}
    check("Every placed Stripe line is a billable Stripe line", set(placed_ids) <= billable_ids)
    skipped_ids = {r["line_item_id"] for r in lines} - billable_ids
    check("No void/uncollectible Stripe line was placed", not (set(placed_ids) & skipped_ids))

    shares = defaultdict(float)
    for r in st_detail:
        shares[r["Line Item ID"]] += r["Amount"]
    net = {r["line_item_id"]: r["net"] for r in billable}
    off = [i for i, v in shares.items() if abs(v - net[i]) > TOL]
    check("Each Stripe line placed once in full (shares add up)", not off, f"{len(off)} lines off" if off else "")

    dupes = table(wb["Duplicate Invoices"])
    unexplained = [d for d in dupes if d["Note"] == "unexplained difference"]
    check("Every removed HubSpot copy matches its Stripe original, or the gap is explained", not unexplained,
          f"{[d['Invoice #'] for d in unexplained[:5]]}" if unexplained else f"{len(dupes)} invoices")
    for note, n in Counter(d["Note"] for d in dupes if d["Note"]).items():
        gap = sum(d["Stripe Invoice Total"] - d["HubSpot Amount (placed)"] for d in dupes if d["Note"] == note)
        print(f"INFO  {n} invoices: {note}; Stripe minus HubSpot = {gap:,.2f}")
    check("No HubSpot invoice removed twice", len({d['HubSpot Invoice ID'] for d in dupes}) == len(dupes))
    check("No Stripe invoice used as the original twice", len({d['Stripe Invoice ID'] for d in dupes}) == len(dupes))
    copies = {d["HubSpot Invoice ID"] for d in dupes}
    native_removed = [r for r in hs_detail if r["Invoice ID"] in copies and r["Invoice Source"] != "data_sync"]
    check("Only Stripe-synced (data_sync) HubSpot invoices were removed", not native_removed)

    dup_cells = sum(g["Synced Duplicates"].values())
    dup_detail = sum(r["Amount"] for r in hs_detail if r["Stripe Copy"] == "yes")
    check("Synced Duplicates grid = HubSpot lines flagged as Stripe copies", abs(dup_cells - dup_detail) < 1,
          f"diff {dup_cells - dup_detail:.2f}")

    print(f"\n{sum(results)}/{len(results)} checks passed")
    return all(results)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    sys.exit(0 if main(*sys.argv[1:]) else 1)
