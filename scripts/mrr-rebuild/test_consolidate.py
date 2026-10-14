"""Unit tests for the MRR rebuild rules (no network). Run: python3 -m unittest test_consolidate"""
import os
import unittest
from datetime import date

os.environ.setdefault("HUBSPOT_ACCESS_TOKEN", "test")

import consolidate_stripe as cs
import populate_fees as hs

FACILITIES = {
    "1": {"name": "Lockwood - Emerald Road Self Storage", "rows": {f: 0 for f in hs.FEE_ROWS}},
    "2": {"name": "Lockwood - Picayune Self Storage", "rows": {f: 0 for f in hs.FEE_ROWS}},
    "3": {"name": "American Storage - Central", "rows": {f: 0 for f in hs.FEE_ROWS}},
    "4": {"name": "Raber's Storage", "rows": {f: 0 for f in hs.FEE_ROWS}},
}


def stripe_row(**kw):
    row = {"customer_name": "Cust", "customer_email": "a@b.com", "customer_id": "cus_1",
           "invoice_number": "INV-1", "invoice_status": "paid", "invoice_date": "2026-03-31",
           "line_item": "Management Fee", "line_amount": "1000.0", "line_discount": "0.0",
           "invoice_total": "1000.0", "currency": "USD", "invoice_id": "in_1", "line_item_id": "il_1"}
    row.update(kw)
    row["net"] = round(float(row["line_amount"]) - float(row["line_discount"]), 2)
    row["invoice_total"] = float(row["invoice_total"])
    return row


def hs_line(**kw):
    ln = {"cid": "3", "facility": "American Storage - Central", "invoice_id": "h1", "invoice_number": "INV-1",
          "invoice_source": "data_sync", "invoice_total": 1000.0, "amount": 1000.0, "fee": "Management Fee",
          "ym": (2026, 4)}
    ln.update(kw)
    return ln


class FeeAndMonthRules(unittest.TestCase):
    def test_category(self):
        cases = {
            "Management Fee - October 2026": "Management Fee",
            "January Mgmt Fee": "Management Fee",
            "Northwest Storage Mgmt Fee ($600 + net tenant protection from prior month)": "Management Fee",
            "Marketing Fee": "Marketing Fee",
            "Anytime Mkting Fee": "Marketing Fee",
            "Bookkeeping": "Bookkeeping Fee",
            "May Books": "Bookkeeping Fee",
            "White Label Set Up Fee": "Setup Fee",
            "White Label Set-up Fee (Half)": "Setup Fee",
            "Website Fee": "Other Fee",
            "Call Center Fee": "Other Fee",
            "Insight PG": "Other Fee",
            "Local Locker Referral Fee": "Other Fee",
            "StorageReach - May 2026": "Reimbursements",
            "AI Lean SaaS - March 2026": "Reimbursements",
            "Sparefoot Reimbursement": "Reimbursements",
            "Marketing Reimbursement- Signage": "Reimbursements",
            "Burkholder Property Management $601.16 — Landscaping": "Reimbursements",
            "Tenant Protection ($1216.43 x 75%)": "Reimbursements",
            "BOG July (Shauna Crowell)": "Reimbursements",
            "Wix Website Hosting": "Reimbursements",
            "Chateau Locks": "Reimbursements",
        }
        for name, fee in cases.items():
            self.assertEqual(hs.category(name), fee, name)

    def test_invoice_month_ignores_month_in_name(self):
        ym, src = hs.invoice_month({"hs_invoice_date": "2026-05-05T12:00:00Z"})
        self.assertEqual((ym, src), ((2026, 5), "invoice date"))

    def test_invoice_date_is_eastern(self):
        # 03:59 UTC on Mar 1 is still Feb 28 in New York
        ym, _ = hs.invoice_month({"hs_invoice_date": "2026-03-01T03:59:59.999Z"})
        self.assertEqual(ym, (2026, 2))

    def test_midnight_eastern_minus_one_ms_is_the_next_day(self):
        # Safe Storage INV-4060 is stored as Apr 1 00:00 EDT minus 1 ms: it is the April invoice
        self.assertEqual(hs.invoice_day({"hs_invoice_date": "2026-04-01T03:59:59.999Z"})[0], date(2026, 4, 1))
        self.assertEqual(hs.invoice_day({"hs_invoice_date": "2026-03-01T04:59:59.999Z"})[0], date(2026, 3, 1))
        self.assertEqual(hs.invoice_day({"hs_invoice_date": "2026-05-06T03:59:59.999Z"})[0], date(2026, 5, 6))

    def test_date_only_value_is_not_shifted(self):
        # HubSpot stores a plain date as midnight UTC: Apr 1 must stay Apr 1, not Mar 31
        ym, _ = hs.invoice_month({"hs_invoice_date": "2026-04-01T00:00:00Z"})
        self.assertEqual(ym, (2026, 4))
        self.assertEqual(hs.invoice_day({"hs_invoice_date": "2026-03-01T00:00:00Z"})[0], date(2026, 3, 1))

    def test_create_date_fallback(self):
        ym, src = hs.invoice_month({"hs_createdate": "2026-04-15T14:00:00Z"})
        self.assertEqual((ym, src), ((2026, 4), "invoice create date"))


class InvoiceAdjustments(unittest.TestCase):
    def test_order_discount_and_late_fee_become_lines(self):
        from unittest import mock
        base = dict(cid="3", facility="F", invoice_number="INV-1", invoice_source="native_invoice",
                    share=1, ym=(2026, 3), invoice_status="paid")
        lines = [dict(base, invoice_id="d", invoice_total=5233.48, line_id="a", fee="Management Fee",
                      full_amount=7000.0, amount=7000.0),
                 dict(base, invoice_id="d", invoice_total=5233.48, line_id="b", fee="Other Fee",
                      full_amount=3233.48, amount=3233.48),
                 dict(base, invoice_id="f", invoice_total=105.0, line_id="c", fee="Management Fee",
                      full_amount=100.0, amount=100.0),
                 dict(base, invoice_id="ok", invoice_total=50.0, line_id="e", fee="Other Fee",
                      full_amount=50.0, amount=50.0)]
        with mock.patch.object(hs, "associations", return_value={}), \
             mock.patch.object(hs, "batch_read", return_value={}):
            adj = hs.invoice_adjustments(lines)
        got = sorted((a["invoice_id"], a["fee"], a["amount"]) for a in adj)
        self.assertEqual(got, [("d", "Management Fee", -5000.0), ("f", "Other Fee", 5.0)])


class DuplicateDetection(unittest.TestCase):
    def test_synced_copy_matches_on_number_and_total(self):
        m = cs.match_synced([hs_line()], [stripe_row()])
        self.assertEqual(m, {"h1": "in_1"})

    def test_native_invoice_with_same_number_is_not_a_duplicate(self):
        m = cs.match_synced([hs_line(invoice_source="native_invoice")], [stripe_row()])
        self.assertEqual(m, {})

    def test_same_number_different_total_is_not_a_duplicate(self):
        m = cs.match_synced([hs_line(invoice_total=999.0)], [stripe_row()])
        self.assertEqual(m, {})


class CustomerMapping(unittest.TestCase):
    def test_narrow_by_name(self):
        got = cs.narrow_by_name("Emerald Road Self Storage (Lockwood-James Davidson)", ["1", "2"], FACILITIES)
        self.assertEqual(got, ["1"])

    def test_narrow_keeps_all_when_no_name_signal(self):
        self.assertEqual(cs.narrow_by_name("Dan Lockwood", ["1", "2"], FACILITIES), ["1", "2"])

    def test_mapping_order(self):
        rows = [stripe_row(customer_id="cus_sync", invoice_id="in_1"),
                stripe_row(customer_id="cus_mail", invoice_id="in_2", customer_email="x@y.com",
                           customer_name="Picayune Self Storage (Lockwood)"),
                stripe_row(customer_id="cus_name", invoice_id="in_3", customer_email="", customer_name="Raber's Storage"),
                stripe_row(customer_id="cus_none", invoice_id="in_4", customer_email="", customer_name="Nobody")]
        inv_cids, cmap = cs.map_customers(rows, [hs_line()], {"h1": "in_1"}, FACILITIES, {"x@y.com": ["1", "2"]})
        self.assertEqual(inv_cids, {"in_1": ["3"]})
        self.assertEqual(cmap["cus_sync"], (["3"], "customer's synced invoices"))
        self.assertEqual(cmap["cus_mail"], (["2"], "contact email"))
        self.assertEqual(cmap["cus_name"], (["4"], "name match"))
        self.assertEqual(cmap["cus_none"], ([], "no match"))


class MonthColumns(unittest.TestCase):
    def test_extend_months_back_to_nov_2022(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Deal Name", "Hubspot ID", "", "", "", "", "", "Jan-23", "Feb-23"])
        ws.append(["Fac", 1, None, None, None, None, None, "=SUM(H3:H6)", "=SUM(I3:I6)"])
        for label in hs.FEE_ROWS:
            ws.append([label])
        hs.extend_months(ws)
        self.assertEqual(sorted(hs.month_columns(ws)), [(2022, 11), (2022, 12), (2023, 1), (2023, 2)])
        self.assertEqual([ws.cell(2, c).value for c in range(8, 12)],
                         ["=SUM(H3:H8)", "=SUM(I3:I8)", "=SUM(J3:J8)", "=SUM(K3:K8)"])

    def test_expand_blocks_adds_setup_and_reimbursement_rows(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Deal Name", "Hubspot ID", "", "", "", "", "", "Jan-23"])
        for cid in (1, 2):
            ws.append([f"Fac {cid}", cid])
            for label in hs.FEE_ROWS[:4]:
                ws.append([label])
        hs.expand_blocks(ws)
        self.assertEqual([ws.cell(r, 1).value for r in range(2, 16)],
                         ["Fac 1", *hs.FEE_ROWS, "Fac 2", *hs.FEE_ROWS])
        self.assertEqual(hs.read_facilities(ws)["2"]["rows"]["Reimbursements"], 15)


    def test_add_company_block_appends_once(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Deal Name", "Hubspot ID", "", "", "", "", "", "Jan-23"])
        ws.append(["Fac", 1])
        for label in hs.FEE_ROWS:
            ws.append([label])
        hs.add_company_blocks(ws, {"99": "Corporate"})
        hs.add_company_blocks(ws, {"99": "Corporate"})
        hs.extend_months(ws, first=(2023, 1))
        self.assertEqual(ws.max_row, 15)
        self.assertEqual((ws.cell(9, 1).value, ws.cell(9, 2).value, ws.cell(9, 8).value),
                         ("Corporate", 99, "=SUM(H10:H15)"))
        self.assertEqual(hs.read_facilities(ws)["99"]["rows"]["Reimbursements"], 15)


class StripeLines(unittest.TestCase):
    def test_skips_void_and_uncollectible_and_uses_net(self):
        rows = [stripe_row(), stripe_row(line_item_id="il_2", invoice_status="void"),
                stripe_row(line_item_id="il_3", invoice_status="uncollectible"),
                stripe_row(line_item_id="il_4", line_discount="100.0")]
        lines, unmatched, skipped = cs.stripe_lines(rows, {"in_1": ["3"]}, {}, FACILITIES)
        self.assertEqual(len(skipped), 2)
        self.assertEqual([ln["amount"] for ln in lines], [1000.0, 900.0])
        self.assertEqual(lines[0]["ym"], (2026, 3))

    def test_split_across_facilities(self):
        lines, _, _ = cs.stripe_lines([stripe_row()], {"in_1": ["1", "2"]}, {}, FACILITIES)
        self.assertEqual([(ln["cid"], ln["amount"]) for ln in lines], [("1", 500.0), ("2", 500.0)])

    def test_unmatched_customer(self):
        lines, unmatched, _ = cs.stripe_lines([stripe_row()], {}, {"cus_1": ([], "no match")}, FACILITIES)
        self.assertEqual((lines, len(unmatched)), ([], 1))

    def test_month_is_invoice_date_not_name(self):
        lines, _, _ = cs.stripe_lines([stripe_row(line_item="Sparefoot Reimbursement - March 2026",
                                                  invoice_date="2026-05-05")], {"in_1": ["3"]}, {}, FACILITIES)
        self.assertEqual((lines[0]["ym"], lines[0]["invoice_date"]), ((2026, 5), date(2026, 5, 5)))


if __name__ == "__main__":
    unittest.main()
