#!/usr/bin/env python3
"""HubSpot Software Audit: compares software fields on Deals and Companies and writes an XLSX.

Inputs are the raw outputs of the HubSpot query_crm_data tool, saved to files (the tool
saves large results to a file automatically; save small inline results to a file yourself).

  --companies FILE [FILE ...]  output(s) of the COMPANY query (hs_object_id, name, previous_software, prior_software)
  --deals FILE [FILE ...]      output(s) of the DEAL query (hs_object_id, dealname, dealstage, pipeline,
                               current_mgmt_software, prior_software, COMPANY.hs_object_id)
  --out PATH                   where to write the workbook
  --expect-companies N / --expect-deals N   counts from COUNT(*) queries, to catch truncated pulls

Prints a JSON summary (counts per check, plus warnings) to stdout.
"""
import argparse, collections, datetime, html, json, re, sys

# ---------- loading ----------

TSV_KEYS = {  # (label, internal name) -> our key; label disambiguates the two hs_object_id columns
    ('Deal', 'hs_object_id'): 'deal_id', ('Company', 'hs_object_id'): 'company_id',
    ('Deal Name', 'dealname'): 'dealname', ('Pipeline', 'pipeline'): 'pipeline',
    ('Deal Stage', 'dealstage'): 'dealstage', ('Prior Software', 'prior_software'): 'prior_software',
    ('Current Mgmt Software', 'current_mgmt_software'): 'current_mgmt_software',
    ('Company name', 'name'): 'name',
    ('Current Software (at onboarding)', 'previous_software'): 'previous_software',
}

def _blank(v):
    v = html.unescape(v or '').strip()
    return None if v in ('', 'Unassigned') else v

def read_records(path, warnings):
    """Return a list of dicts from a saved query_crm_data result (JSON-record or TSV shape)."""
    raw = open(path, encoding='utf-8').read()
    try:
        contents = [r['content'] for r in json.loads(raw)['results']]
    except (ValueError, KeyError, TypeError):
        contents = [raw]  # plain TSV pasted into a file
    out = []
    for c in contents:
        if 'Dataset TSV:' in c or '\t' in c.split('\n', 1)[0]:
            body = c.split('Dataset TSV:\n', 1)[-1]
            m = re.search(r'\n\nShowing (\d+) of (\d+) total results', body)
            if m:
                body = body[:m.start()]
                if m.group(1) != m.group(2):
                    warnings.append(f'{path}: only {m.group(1)} of {m.group(2)} rows returned - pull the rest and pass both files')
            lines = body.strip('\n').split('\n')
            hdr = []
            for h in lines[0].split('\t'):
                mm = re.match(r'(.*?) \[(.*?)\]$', h.strip())
                label, internal = (mm.group(1), mm.group(2)) if mm else (h, h)
                hdr.append(TSV_KEYS.get((label, internal), internal))
            for l in lines[1:]:
                f = l.split('\t')
                if len(f) != len(hdr):
                    warnings.append(f'{path}: skipped malformed row: {l[:80]}')
                    continue
                out.append({k: _blank(v) for k, v in zip(hdr, f)})
        else:
            try:
                p = json.loads(c).get('properties', {})
            except ValueError:
                continue
            out.append({k: _blank(v) for k, v in p.items()})
    return out

def strip_ids(v):
    return re.sub(r' \(\d+\)$', '', v or '')

# ---------- normalising free text to the Prior Software dropdown ----------

VENDORS = [
    ('Hummingbird', r'hummingbird'),
    ('StorEDGE', r'stor\s*e?\s*edge|storeedge|storab[el]{1,2}e?[\s/-]*edge|storabe edge|^edge\b|storedgefms'),
    ('Sitelink', r'site\s*lin[ke]'),
    ('Easy Storage Solutions (ESS)', r'\bess\b|eas[ey] (self )?storage|ez storage|storab[el]{1,2}e? ?eas[ty]|storable ez\b|easy storable|^easy$|storageunitsoftware'),
    ('Web Self Storage', r'web\s*-?\s*self[\s-]*storage|\bwss\b|\bwws\b|web storage'),
    ('Tenant Inc', r'ten+[ae]nt'),
    ('Cubby', r'cubby'),
    ('CC Storage', r'cc storage|^cc$'),
    ('QuikStor', r'qui?c?k\s*stor\b|quikstor'),
    ('Yardi', r'yardi'),
]
UNKNOWN = (r"^we$|^stora$|isn'?t sure|unknown|don'?t know|doesn'?t know|didn'?t know|not sure|unsure|^\?$|^tbd$|"
           r"unconfirmed|^software$|^storable$|storable \(not sure|storable, not sure|something from storable|"
           r"some sort of storable|same as|has info")
NONE = (r"^(none|na|n/a|n\\a|nothing|new|new build|0|excel|uses excel|excel lol|excel spreadsheets)\b|^none|^n/?a\b|"
        r"nothing yet|pen and paper|manages manually|self managed|backs of envelopes|under construction|new facility")
FUTURE = r"\b(will be|moving to|may switch|switching to|but will switch|transitioning|previously|switched from|but storedge will)\b.*$"

def interpret(v):
    """Map free text to a dropdown value. Returns (value, status) with status OK/Multiple/Unknown/Blank."""
    if not v:
        return '', 'Blank'
    s = re.sub(r'^previous:\s*', '', v.strip().lower())
    core = re.sub(FUTURE, '', s).strip(' ,.-(')  # compare what is in use now, not future plans
    if not core:
        return '', 'Unknown'
    found = [n for n, p in VENDORS if re.search(p, core)]
    if 'Hummingbird' in found and 'Tenant Inc' in found and not re.search(r'&|and (?!.*hummingbird)', core):
        found.remove('Tenant Inc')  # "Hummingbird by Tenant Inc" is one product
    if len(found) == 1:
        return found[0], 'OK'
    if found:
        return ' / '.join(found), 'Multiple'
    if re.search(UNKNOWN, core):
        return '', 'Unknown'
    if re.search(NONE, core):
        return 'None', 'OK'
    return 'Other', 'OK'

def redact(v):
    if v and re.search(r'@|\bpw\b|password|passwd|login:', v, re.I):
        return '[REDACTED - field contains login details]'
    return v or ''

def compare(text, enum):
    cat, st = interpret(text)
    if not text and not enum: return cat, 'Both blank'
    if not text: return cat, 'Text field blank'
    if not enum: return cat, 'Prior Software blank'
    if st == 'Unknown': return cat, 'Review - text unclear'
    if st == 'Multiple':
        return cat, 'Review - multiple (includes prior)' if enum in cat.split(' / ') else 'Review - multiple softwares'
    if cat == enum: return cat, 'Match'
    if {cat, enum} == {'Tenant Inc', 'Hummingbird'}: return cat, 'Close - Tenant Inc vs Hummingbird'
    return cat, 'Mismatch'

# ---------- audit ----------

def run(company_files, deal_files, expect_c, expect_d):
    warnings = []
    C = {}
    for f in company_files:
        for r in read_records(f, warnings):
            if r.get('hs_object_id'):
                C[r['hs_object_id']] = r
    rows = []
    for f in deal_files:
        rows += read_records(f, warnings)
    rows = [r for r in rows if r.get('deal_id') or r.get('hs_object_id')]
    for r in rows:
        r.setdefault('deal_id', r.get('hs_object_id'))
    seen = set(); rows = [r for r in rows if not ((r['deal_id'], r.get('company_id')) in seen or seen.add((r['deal_id'], r.get('company_id'))))]
    D = {}
    for r in rows:
        D.setdefault(r['deal_id'], dict(r, companies=[]))['companies'].append(r.get('company_id'))
    if expect_c is not None and len(C) != expect_c:
        warnings.append(f'Expected {expect_c} companies but loaded {len(C)} - the pull may be incomplete')
    if expect_d is not None and len(D) != expect_d:
        warnings.append(f'Expected {expect_d} deals but loaded {len(D)} - the pull may be incomplete')
    if rows and 'company_id' not in rows[0]:
        warnings.append('Deal file has no Company ID column - check 3 cannot run; include COMPANY.hs_object_id in the deal query')

    co = []
    for cid, c in C.items():
        cat, res = compare(c.get('previous_software'), c.get('prior_software'))
        co.append(dict(id=cid, name=c.get('name') or '', text=redact(c.get('previous_software')), cat=cat,
                       prior=c.get('prior_software') or '', result=res))
    dl = []
    for did, d in D.items():
        cat, res = compare(d.get('current_mgmt_software'), d.get('prior_software'))
        dl.append(dict(id=did, name=d.get('dealname') or '', pipeline=strip_ids(d.get('pipeline')), stage=strip_ids(d.get('dealstage')),
                       text=redact(d.get('current_mgmt_software')), cat=cat, prior=d.get('prior_software') or '', result=res))
    x = []
    for r in rows:
        cid = r.get('company_id')
        if not cid:
            continue
        c = C.get(cid, {})  # a company missing from the company pull has both company fields blank
        dp, cp = r.get('prior_software'), c.get('prior_software')
        if not dp and not cp: res = 'Both blank'
        elif not cp: res = 'Company Prior Software blank'
        elif not dp: res = 'Deal Prior Software blank'
        else: res = 'Match' if dp == cp else 'Mismatch'
        x.append(dict(cid=cid, cname=c.get('name') or '', cprior=cp or '', did=r['deal_id'], dname=r.get('dealname') or '',
                      pipeline=strip_ids(r.get('pipeline')), stage=strip_ids(r.get('dealstage')), dprior=dp or '', result=res))
    no_co = sum(1 for d in D.values() if not any(d['companies']))
    redacted = [('Deal', r['id'], r['name']) for r in dl if r['text'].startswith('[REDACTED')] + \
               [('Company', r['id'], r['name']) for r in co if r['text'].startswith('[REDACTED')]
    return co, dl, x, no_co, redacted, warnings

# ---------- workbook ----------

ORDER = ['Mismatch', 'Close', 'Review', 'Prior Software blank', 'Text field blank', 'Company Prior Software blank',
         'Deal Prior Software blank', 'Both blank', 'Match']
def rank(r): return next((i for i, k in enumerate(ORDER) if r.startswith(k)), 99)

def write_xlsx(path, co, dl, x, no_co, redacted, run_date):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    colors = [('Match', 'D9EAD3'), ('Mismatch', 'F4CCCC'), ('Close', 'FFF2CC'), ('Review', 'FFF2CC'), ('blank', 'EDEDED')]
    def fill(r):
        for k, v in colors:
            if r.startswith(k) or k in r:
                return PatternFill('solid', fgColor=v)
    wb = Workbook(); ws = wb.active; ws.title = 'Summary'

    def sheet(title, cols, rows, keys):
        s = wb.create_sheet(title); s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864')
        for r in sorted(rows, key=lambda r: (rank(r['result']), r.get('name') or r.get('cname') or '')):
            s.append([r[k] for k in keys])
            f = fill(r['result'])
            if f: s.cell(s.max_row, len(keys)).fill = f
        for i, c in enumerate(cols, 1):
            w = max([len(str(c))] + [len(str(s.cell(j, i).value or '')) for j in range(2, s.max_row + 1)])
            s.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 55)
        s.freeze_panes = 'A2'; s.auto_filter.ref = s.dimensions

    sheet('1 Deal Check', ['Deal ID', 'Deal Name', 'Pipeline', 'Stage', 'Current Mgmt Software (text)', 'Interpreted As', 'Prior Software', 'Result'],
          dl, ['id', 'name', 'pipeline', 'stage', 'text', 'cat', 'prior', 'result'])
    sheet('2 Company Check', ['Company ID', 'Company Name', 'Current Software (at onboarding) (text)', 'Interpreted As', 'Prior Software', 'Result'],
          co, ['id', 'name', 'text', 'cat', 'prior', 'result'])
    sheet('3 Company vs Deal', ['Company ID', 'Company Name', 'Company Prior Software', 'Deal ID', 'Deal Name', 'Pipeline', 'Stage', 'Deal Prior Software', 'Result'],
          x, ['cid', 'cname', 'cprior', 'did', 'dname', 'pipeline', 'stage', 'dprior', 'result'])

    ws.append([f'HubSpot Software Audit - {run_date}']); ws['A1'].font = Font(bold=True, size=14); ws.append([])
    def block(title, rows, note):
        ws.append([title]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append([note]); ws.cell(ws.max_row, 1).font = Font(italic=True, color='666666')
        for k, v in sorted(collections.Counter(r['result'] for r in rows).items(), key=lambda kv: rank(kv[0])):
            ws.append([k, v]); f = fill(k)
            if f: ws.cell(ws.max_row, 1).fill = f
        ws.append(['Total', len(rows)]); ws.cell(ws.max_row, 1).font = Font(bold=True); ws.append([])
    block('1. Deal: Current Mgmt Software vs Prior Software', dl, 'All deals where either field is filled in')
    block('2. Company: Current Software (at onboarding) vs Prior Software', co, 'All companies where either field is filled in')
    block('3. Company Prior Software vs associated Deal Prior Software', x,
          f'One row per deal-company pair. {no_co} deals with software data have no associated company and are not in this check.')
    if redacted:
        ws.append([f'{len(redacted)} records have logins/passwords typed into a software field (redacted in this file) - move them to the password manager and clear them in HubSpot:'])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color='C00000')
        for t, i, n in redacted:
            ws.append([f'{t} {i} - {n}'])
        ws.append([])
    ws.append(['How text was matched']); ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for t in ['The text fields are free text, so each value is mapped to a Prior Software dropdown option before comparing.',
              'e.g. "storedge", "Storable Edge" -> StorEDGE; "ESS", "Storable Easy" -> Easy Storage Solutions (ESS); "WSS" -> Web Self Storage; "na", "Excel", "new build" -> None.',
              'Software not in the dropdown (Storage Commander, Syrasoft, 6Storage, CubeSmart, etc.) -> Other.',
              'Future plans ("will be StorEdge", "moving to StorEDGE") are ignored; the software in use now is compared.',
              '"Unknown", bare "Storable" and values naming several softwares are marked Review rather than Match/Mismatch.',
              'Tenant Inc vs Hummingbird (same vendor) is marked Close rather than Mismatch.']:
        ws.append([t])
    ws.column_dimensions['A'].width = 70; ws.column_dimensions['B'].width = 10
    wb.save(path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--companies', nargs='+', required=True)
    ap.add_argument('--deals', nargs='+', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--expect-companies', type=int)
    ap.add_argument('--expect-deals', type=int)
    ap.add_argument('--date', default=datetime.date.today().isoformat())
    a = ap.parse_args()
    co, dl, x, no_co, redacted, warnings = run(a.companies, a.deals, a.expect_companies, a.expect_deals)
    write_xlsx(a.out, co, dl, x, no_co, redacted, a.date)
    def counts(rows): return dict(sorted(collections.Counter(r['result'] for r in rows).items(), key=lambda kv: rank(kv[0])))
    def by_pipeline(rows):
        out = collections.defaultdict(collections.Counter)
        for r in rows: out[r['pipeline']][r['result']] += 1
        return {k: dict(v) for k, v in out.items()}
    json.dump(dict(output=a.out, companies=len(co), deals=len(dl), deal_company_pairs=len(x), deals_without_company=no_co,
                   deal_check=counts(dl), deal_check_by_pipeline=by_pipeline(dl), company_check=counts(co), company_vs_deal=counts(x),
                   records_with_logins=[f'{t} {i} - {n}' for t, i, n in redacted], warnings=warnings), sys.stdout, indent=1)

if __name__ == '__main__':
    main()
