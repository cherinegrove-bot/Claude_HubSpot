#!/usr/bin/env python3
"""Build the one-page HTML summary of a full-map audit from the finished workbook, findings.json and the
data fetch.py already saved. Nothing is fetched from HubSpot.

Usage: python3 summary_html.py --workbook audits/x/WLS_X_Workflow_Audit.xlsx --findings audits/x/findings.json \
                               --work $WK --team "CS Ops" --out audits/x/cs-ops-audit-summary.html
findings.json may contain  "top": [{"id": "F-03", "text": "plain-English line", "next": "what to check"}]
and, on reviewed findings, "workflow_ids": ["123", ...] to link them to workflows.
"""
import argparse, collections, datetime as dt, html, json, os, re, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rules

ap = argparse.ArgumentParser()
ap.add_argument('--workbook', required=True)
ap.add_argument('--findings')
ap.add_argument('--work', required=True)
ap.add_argument('--team', required=True)
ap.add_argument('--out', required=True)
A = ap.parse_args()
W = lambda n, d=None: json.load(open(os.path.join(A.work, n))) if os.path.exists(os.path.join(A.work, n)) else d
E = lambda s: html.escape(str(s if s is not None else ''))

wb = openpyxl.load_workbook(A.workbook, read_only=True)
AF = [r for r in wb['Audit Findings'].iter_rows(min_row=2, values_only=True) if r and r[0]]
INV = [r for r in wb['Workflow Inventory'].iter_rows(min_row=2, values_only=True) if r and r[0]]
EXTRA = json.load(open(A.findings)) if A.findings and os.path.exists(A.findings) else {}
CFG = W('configs.json', {})
TASKS = W('tasks.json', [])
SCOPE = W('scope.json', {})
PROPS = W('properties.json', {})
LIST_NAMES = W('list_names.json', {})
TEAM_NAME, TEAM = rules.load_team(A.team)
QIDS = SCOPE.get('queue_ids') or (TEAM or {}).get('queue_ids') or []
SINCE = SCOPE.get('since') or (TEAM or {}).get('start_date')
RUN = dt.datetime.fromtimestamp(os.path.getmtime(os.path.join(A.work, 'COMPLETE')) if os.path.exists(os.path.join(A.work, 'COMPLETE')) else os.path.getmtime(A.workbook)).strftime('%Y-%m-%d')


# ------------------------------------------------------------------ plain-English trigger
def label(prop):
    for obj in ('companies', 'tickets', 'deals', 'contacts', 'tasks'):
        if prop in PROPS.get(obj, {}):
            return PROPS[obj][prop]['label']
    return prop


def value(prop, v):
    for obj in ('companies', 'tickets', 'deals', 'contacts', 'tasks'):
        o = PROPS.get(obj, {}).get(prop, {}).get('options', {})
        if str(v) in o:
            return o[str(v)]
    return v


def cond(f):
    if f.get('filterType') == 'IN_LIST' or 'listId' in f:
        nm = LIST_NAMES.get(str(f.get('listId')))
        return f'on list "{nm}"' if nm else f'on list {f.get("listId")}'
    op = f.get('operation', {})
    o = op.get('operator', '')
    vals = op.get('values') or ([op['value']] if 'value' in op else [])
    p = label(f.get('property'))
    if o == 'IS_KNOWN':
        return f'{p} is set'
    if o == 'IS_UNKNOWN':
        return f'{p} is empty'
    if o in ('IS_ANY_OF', 'IS_EQUAL_TO'):
        return f'{p} is {" or ".join(str(value(f["property"], v)) for v in vals)}'
    if o in ('IS_NONE_OF', 'IS_NOT_EQUAL_TO'):
        return f'{p} is not {" or ".join(str(value(f["property"], v)) for v in vals)}'
    if o == 'IS_BETWEEN' and op.get('lowerBoundTimePoint', {}).get('offset'):
        return f'{p} within the last {abs(op["lowerBoundTimePoint"]["offset"].get("days", 0))} days'
    if 'TIME' in str(op.get('operationType')):
        return f'{p} ({o.lower().replace("_", " ")} a date)'
    return f'{p} {o.lower().replace("_", " ")} {", ".join(map(str, vals))}'.strip()


def branch_text(fb):
    parts = [cond(f) for f in fb.get('filters', [])]
    for b in fb.get('filterBranches', []):
        if b.get('filterBranchType') == 'ASSOCIATION':
            parts.append('has an associated ' + {'0-49': 'email', '0-2': 'company', '0-5': 'ticket', '0-1': 'contact'}.get(b.get('objectTypeId'), 'record') + ' where ' + branch_text(b))
        else:
            t = branch_text(b)
            if t:
                parts.append(t)
    return (' or ' if fb.get('filterBranchType') == 'OR' else ' and ').join(p for p in parts if p)


def trigger_plain(c):
    e = c.get('enrollmentCriteria', {})
    obj = {'0-2': 'Company', '0-5': 'Ticket', '0-3': 'Deal', '0-1': 'Contact', '0-27': 'Task'}.get(c.get('objectTypeId'), 'Record')
    if e.get('type') == 'LIST_BASED':
        t = f'{obj} where {branch_text(e.get("listFilterBranch", {}))}'
    elif e.get('type') == 'MANUAL':
        t = 'Manual enrollment only'
    else:
        t = f'{obj} event trigger'
    sb = rules.suppression(c)
    if sb:
        t += f'; not if {branch_text(sb)}'
    sc = rules.schedule_text(c)
    if sc:
        t += f'; {sc}'
    return t


# ------------------------------------------------------------------ numbers
by_class = collections.Counter(r[1] for r in AF)
on = sum(1 for c in CFG.values() if c.get('isEnabled'))
off = len(CFG) - on
in_scope_tasks = [t for t in TASKS if t.get('_flow') in CFG and t['properties'].get('hs_task_is_sub_task') != 'true' and (not SINCE or t['properties']['hs_createdate'][:10] >= SINCE)]
tasks_by_flow = collections.Counter(t['_flow'] for t in in_scope_tasks)

# finding <-> workflow links
wf_ids_by_text = {}
for m in EXTRA.get('findings', []):
    if 'workflow_ids' in m:
        wf_ids_by_text[m['finding'][:80]] = set(map(str, m['workflow_ids']))
short = {w: re.sub(r'^create tasks \|\s*', '', c['name'].strip().lower()) for w, c in CFG.items()}


def mentions(hay, name):
    """Exact mention: the name must not continue into a longer workflow name."""
    return bool(re.search(re.escape(name) + r'(?=$|[,.;:)"\']|\s\(|\s(?:and|or|has|is|-|x\d|\d))', hay))


related = collections.defaultdict(list)
for r in AF:
    fid, wfcol, text = r[0], str(r[2] or ''), str(r[4] or '')
    ids = wf_ids_by_text.get(text[:80])
    if ids is None and re.match(r'\s*(all\b|all in scope|several|team queue)', wfcol.lower()):
        continue  # portfolio-wide findings are not linked to each workflow
    hay = (wfcol + ' ' + text).lower()
    for w, c in CFG.items():
        if ids is not None:
            if w in ids:
                related[w].append(fid)
        elif mentions(hay, c['name'].strip().lower()) or (len(short[w]) > 12 and mentions(hay, short[w])) or w in hay:
            related[w].append(fid)

CLS = {'Confirmed': ('#E2EFDA', '#3B7D23'), 'Potential Issue': ('#FFF2CC', '#B7860B'), 'Needs Verification': ('#FCE4D6', '#C55A11'), 'No Issue Found': ('#DDEBF7', '#2E75B6')}
af_by_id = {r[0]: r for r in AF}

# ------------------------------------------------------------------ HTML
css = """
*{box-sizing:border-box}body{margin:0;font-family:Poppins,system-ui,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;color:#1f2933;background:#fff;font-size:14px;line-height:1.5}
.bar{background:#263449;color:#fff;padding:22px 28px}.bar h1{margin:0 0 6px;font-size:22px;font-weight:600}.meta{display:flex;flex-wrap:wrap;gap:6px 22px;font-size:13px;opacity:.9}
main{max-width:1180px;margin:0 auto;padding:22px 28px 40px}h2{font-size:17px;margin:30px 0 12px;color:#263449;font-weight:600}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}.card{border:1px solid #dfe3e8;border-radius:8px;padding:12px 14px;background:#F3F5F7}
.card .n{font-size:26px;font-weight:600;color:#263449}.card .l{font-size:12px;color:#52606d}.card.c{border-left:5px solid var(--a);background:var(--b)}
.top{display:grid;gap:10px}.tf{border:1px solid #dfe3e8;border-left:5px solid var(--a);border-radius:8px;padding:12px 14px;background:#fff}
.tf .h{display:flex;gap:10px;align-items:baseline;flex-wrap:wrap}.id{font-weight:600;color:#263449}.pill{font-size:11px;padding:2px 8px;border-radius:10px;background:var(--b);color:var(--a);font-weight:600}
.tf p{margin:6px 0 0}.next{color:#52606d;font-size:13px}.next b{color:#3F5066}
.tw{overflow-x:auto;border:1px solid #dfe3e8;border-radius:8px}table{border-collapse:collapse;width:100%;font-size:13px}
th{background:#3F5066;color:#fff;text-align:left;padding:8px 10px;font-weight:600;position:sticky;top:0}td{padding:8px 10px;border-top:1px solid #e4e8ec;vertical-align:top}
tbody tr:nth-child(even) td{background:#F3F5F7}td.cls{border-left:5px solid var(--a);white-space:nowrap}td.fid{white-space:nowrap}.on{color:#3B7D23;font-weight:600}.off{color:#8a94a0;font-weight:600}
footer{margin-top:34px;padding:16px 18px;background:#F3F5F7;border-radius:8px;font-size:12.5px;color:#3e4c59}footer b{color:#263449}
@media (max-width:640px){.bar,main{padding-left:16px;padding-right:16px}}
"""
out = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
       f'<title>{E(TEAM_NAME)} Workflow Audit</title><style>{css}</style></head><body>']
out.append(f'<div class="bar"><h1>WLS {E(TEAM_NAME)} Workflow Audit</h1><div class="meta">'
           f'<span>Queue: {E((TEAM or {}).get("queue_name", ""))}{" (" + E(", ".join(QIDS)) + ")" if QIDS else ""}</span>'
           f'<span>Start date: {E(SINCE or "none")}</span><span>Run: {E(RUN)}</span><span>Workflows checked: {len(CFG)}</span></div></div><main>')
# headline numbers
out.append('<h2>Headline numbers</h2><div class="cards">')
for k, (bg, ac) in CLS.items():
    out.append(f'<div class="card c" style="--a:{ac};--b:{bg}"><div class="n">{by_class.get(k, 0)}</div><div class="l">{E(k)}</div></div>')
out.append(f'<div class="card"><div class="n">{on} / {off}</div><div class="l">Workflows ON / OFF</div></div>'
           f'<div class="card"><div class="n">{len(in_scope_tasks)}</div><div class="l">Tasks checked in the queue since {E(SINCE or "start")}</div></div></div>')
# top findings
tops = EXTRA.get('top', [])
if tops:
    out.append('<h2>Top findings</h2><div class="top">')
    for t in tops:
        r = af_by_id.get(t['id'])
        cls = r[1] if r else ''
        bg, ac = CLS.get(cls, ('#F3F5F7', '#3F5066'))
        out.append(f'<div class="tf" style="--a:{ac};--b:{bg}"><div class="h"><span class="id">{E(t["id"])}</span><span class="pill">{E(cls)}</span>'
                   f'<span>{E(r[2] if r else "")}</span></div><p>{E(t["text"])}</p><p class="next"><b>Check next:</b> {E(t.get("next", r[7] if r else ""))}</p></div>')
    out.append('</div>')
# all findings
out.append('<h2>All findings</h2><div class="tw"><table><thead><tr><th>#</th><th>Classification</th><th>Workflow</th><th>Finding</th><th>Recommended verification</th></tr></thead><tbody>')
for r in AF:
    bg, ac = CLS.get(r[1], ('#fff', '#3F5066'))
    out.append(f'<tr><td class="fid"><b>{E(r[0])}</b></td><td class="cls" style="--a:{ac};background:{bg}">{E(r[1])}</td><td>{E(r[2])}</td><td>{E(r[4])}</td><td>{E(r[7])}</td></tr>')
out.append('</tbody></table></div>')
# workflows
out.append('<h2>Workflows</h2><div class="tw"><table><thead><tr><th>Workflow</th><th>ID</th><th>Status</th><th>Trigger (plain English)</th><th>Tasks since start</th><th>Related findings</th></tr></thead><tbody>')
for r in INV:
    w = str(r[1])
    c = CFG.get(w, {})
    st = 'ON' if c.get('isEnabled') else 'OFF'
    out.append(f'<tr><td>{E(r[0])}</td><td>{E(w)}</td><td class="{st.lower()}">{st}</td><td>{E(trigger_plain(c) if c else r[4])}</td>'
               f'<td>{tasks_by_flow.get(w, 0)}</td><td>{E(", ".join(sorted(set(related.get(w, [])), key=lambda x: int(x[2:]))) or "-")}</td></tr>')
for n in EXTRA.get('not_readable', []):
    out.append(f'<tr><td>{E(n.get("name"))}</td><td>-</td><td class="off">NOT FOUND</td><td>{E(n.get("note", ""))}</td><td>-</td><td>{E(n.get("finding", "-"))}</td></tr>')
out.append('</tbody></table></div>')
out.append('<footer><p><b>Evidence labels:</b> VERIFIED (config) = read from the workflow configuration (HubSpot Automation API) · VERIFIED (records) = read from HubSpot task / company / ticket records · '
           'VERIFIED (screenshot) = seen in a screenshot supplied by WLS · INFERENCE = derived, not directly visible · NEEDS VERIFICATION = cannot be determined. Full evidence for every finding is in the workbook.</p>'
           '<p><b>Read only: nothing in HubSpot was changed.</b></p></footer></main></body></html>')
os.makedirs(os.path.dirname(os.path.abspath(A.out)), exist_ok=True)
open(A.out, 'w').write('\n'.join(out))
print(f'{A.out}: {len(AF)} findings, {len(INV)} workflows, {len(in_scope_tasks)} tasks since {SINCE}, {len(tops)} top findings')
