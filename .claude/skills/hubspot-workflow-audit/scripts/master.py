#!/usr/bin/env python3
"""Master rules file: build it once, read it, and compare live workflows with it.

The master rules file (audits/<team>-workflow-audit/<Team>_Master_Rules.xlsx) records how every
workflow of a team should work. It is the source of truth for "Did anyone change the workflows?".
NEVER rebuild it from live workflows during an audit; update a row only when the user confirms
the change was planned.

Usage:
  build (first time, from a saved settings snapshot):
    python3 master.py build --team cs-ops --work <snapshot dir> [--purposes purposes.json] [--source-note "..."]
  update rows the user confirmed as planned changes:
    python3 master.py update --team cs-ops --work <weekly work dir> --ids 123 456 --reason "confirmed by ... on ..."
  remove rows for workflows the user confirmed were deleted:
    python3 master.py remove --team cs-ops --work <weekly work dir> --ids 123 --reason "deleted on purpose, confirmed by ... on ..."
  show differences (no file is written):
    python3 master.py diff --team cs-ops --work <weekly work dir>
"""
import argparse, datetime as dt, difflib, hashlib, json, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rules, describe

COLUMNS = ['Workflow Name', 'Workflow ID', 'Object', 'On/Off', 'Purpose', 'Trigger', 'Schedule', 'Re-enrollment', 'Suppression',
           'Branches', 'Task actions', 'Other actions', 'Number of actions', 'Key actions', 'Task types it creates', 'Queue',
           'Revision', 'Last updated', 'Notes']
# Columns compared by the change check (Question 3), in report order.
COMPARED = ['Workflow Name', 'On/Off', 'Trigger', 'Schedule', 'Re-enrollment', 'Suppression', 'Branches', 'Task actions',
            'Other actions', 'Queue', 'Revision']
MULTILINE = {'Branches', 'Task actions', 'Other actions'}
KEEP = ['id', 'name', 'isEnabled', 'objectTypeId', 'revisionId', 'updatedAt', 'startActionId', 'enrollmentCriteria',
        'enrollmentSchedule', 'suppressionFilterBranch', 'timeWindows', 'blockedDates', 'actions']
CHUNK = 30000
ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', '..', '..'))


# ------------------------------------------------------------------ team files
def team_dir(team):
    return os.path.join(ROOT, 'audits', f'{team}-workflow-audit')


def load_team(team):
    """Settings block (```json ... ```) from audits/<team>-workflow-audit/team-rules.md."""
    path = os.path.join(team_dir(team), 'team-rules.md')
    if not os.path.exists(path):
        sys.exit(f'STOP: {path} not found. Set up the team first (SKILL.md, "Set up a new team").')
    m = re.search(r'```json\s*\n(.*?)\n```', open(path).read(), re.S)
    if not m:
        sys.exit(f'STOP: no ```json settings block in {path}.')
    cfg = json.loads(m.group(1))
    cfg['_dir'] = team_dir(team)
    return cfg


def master_path(team_cfg):
    return os.path.join(team_cfg['_dir'], team_cfg['master_rules_file'])


# ------------------------------------------------------------------ one workflow -> row
def trim(cfg):
    """The parts of a workflow's settings the audit needs. Task descriptions are replaced by a fingerprint."""
    c = {k: cfg.get(k) for k in KEEP if k in cfg}
    acts = []
    for a in cfg.get('actions', []):
        a = json.loads(json.dumps(a))
        f = a.get('fields', {})
        if isinstance(f.get('body'), str):
            f['body'] = {'_sha1': hashlib.sha1(f['body'].encode()).hexdigest()[:12], '_len': len(describe.strip(f['body']))}
        acts.append(a)
    c['actions'] = acts
    return c


def _body_text(cfg):
    """Task description fingerprints, so a changed description shows as a change."""
    out = []
    for a, br in describe.paths(cfg):
        if describe.akind(a) == '0-3':
            b = a.get('fields', {}).get('body')
            fp = b['_sha1'] if isinstance(b, dict) else hashlib.sha1((b or '').encode()).hexdigest()[:12]
            out.append(fp)
    return out


def task_types_for(cfg, team_cfg):
    titles = [(a.get('fields', {}).get('subject') or '') for a in cfg.get('actions', []) if describe.akind(a) == '0-3']
    found, unmatched = [], []
    for t in titles:
        hit = [tt['name'] for tt in team_cfg.get('task_types', []) if str(cfg.get('id')) in tt['workflows'] and re.search(tt['title'], rules.norm(t), re.I)]
        for h in hit:
            if h not in found:
                found.append(h)
        if not hit and t.strip() not in unmatched:
            unmatched.append(t.strip())
    if str(cfg.get('id')) in team_cfg.get('expected_off', []):
        return 'None (switched off as expected)' + (f'; titles: {", ".join(unmatched)}' if unmatched else '')
    if not titles:
        return 'None (creates no tasks)'
    return '; '.join(found) + (f'; NOT A KNOWN TASK TYPE: {", ".join(unmatched)}' if unmatched else '')


def queue_text(cfg, team_cfg):
    qs = [str(a.get('fields', {}).get('queue_id') or '') for a in cfg.get('actions', []) if describe.akind(a) == '0-3']
    if not qs:
        return 'n/a (no task actions)'
    name = lambda q: f'{team_cfg.get("queue_name")} ({q})' if q in team_cfg.get('queue_ids', []) else (q or 'no queue')
    distinct = sorted(set(qs))
    if len(distinct) == 1:
        return f'{name(distinct[0])} on all {len(qs)} task action(s)'
    return '; '.join(f'{name(q)} on {qs.count(q)} of {len(qs)}' for q in distinct)


def key_actions(cfg):
    titles = []
    for a in cfg.get('actions', []):
        if describe.akind(a) == '0-3':
            t = (a.get('fields', {}).get('subject') or '').strip()
            if t and rules.norm(t) not in [rules.norm(x) for x in titles]:
                titles.append(t)
    props = []
    for a in cfg.get('actions', []):
        if describe.akind(a) == 'LIST_BRANCH':
            for p in re.findall(r'"property": "(\w+)"', json.dumps(a)):
                lbl = describe.plabel(p)
                if lbl not in props:
                    props.append(lbl)
    out = []
    if titles:
        out.append('Creates task(s): ' + '; '.join(f'"{t}"' for t in titles))
    if props:
        out.append('Branches on: ' + ', '.join(props))
    return '. '.join(out) or 'No actions that create or change anything'


def row(cfg, team_cfg, purpose='', notes=''):
    return {
        'Workflow Name': cfg.get('name'),
        'Workflow ID': str(cfg.get('id')),
        'Object': describe.OBJ.get(cfg.get('objectTypeId'), cfg.get('objectTypeId')),
        'On/Off': 'ON' if cfg.get('isEnabled') else 'OFF',
        'Purpose': purpose,
        'Trigger': describe.trigger_text(cfg),
        'Schedule': describe.schedule_text(cfg),
        'Re-enrollment': describe.reenroll_text(cfg),
        'Suppression': describe.suppression_text(cfg),
        'Branches': '\n'.join(describe.branch_lines(cfg)) or 'None',
        'Task actions': '\n'.join(f'{l} | description id {fp}' for l, fp in zip(describe.task_lines(cfg, team_cfg.get('queue_name')), _body_text(cfg))) or 'None',
        'Other actions': '\n'.join(describe.other_lines(cfg)) or 'None',
        'Number of actions': describe.action_count(cfg),
        'Key actions': key_actions(cfg),
        'Task types it creates': task_types_for(cfg, team_cfg),
        'Queue': queue_text(cfg, team_cfg),
        'Revision': str(cfg.get('revisionId')),
        'Last updated': (cfg.get('updatedAt') or '')[:19].replace('T', ' ') + ' UTC',
        'Notes': notes,
    }


# ------------------------------------------------------------------ workbook
def _styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style='thin', color='BFBFBF')
    return dict(hfont=Font(name='Arial', bold=True, color='FFFFFF', size=10), hfill=PatternFill('solid', fgColor='263449'),
                font=Font(name='Arial', size=10), bold=Font(name='Arial', size=10, bold=True), title=Font(name='Arial', size=14, bold=True, color='263449'),
                wrap=Alignment(wrap_text=True, vertical='top'), border=Border(left=thin, right=thin, top=thin, bottom=thin))


def write_workbook(path, team_cfg, rows, machine, readme):
    from openpyxl import Workbook
    S = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Read Me'
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 110
    ws.append([f'{team_cfg["team_name"]} master rules'])
    ws['A1'].font = S['title']
    ws.append([])
    for k, v in readme:
        ws.append([k, v])
        r = ws.max_row
        ws.cell(r, 1).font = S['bold']
        ws.cell(r, 2).font = S['font']
        ws.cell(r, 1).alignment = ws.cell(r, 2).alignment = S['wrap']

    ws = wb.create_sheet('Master Rules')
    widths = {'Workflow Name': 40, 'Workflow ID': 13, 'Object': 10, 'On/Off': 8, 'Purpose': 50, 'Trigger': 60, 'Schedule': 28, 'Re-enrollment': 28,
              'Suppression': 36, 'Branches': 90, 'Task actions': 110, 'Other actions': 50, 'Number of actions': 26, 'Key actions': 60,
              'Task types it creates': 40, 'Queue': 30, 'Revision': 10, 'Last updated': 20, 'Notes': 40}
    ws.append(COLUMNS)
    from openpyxl.utils import get_column_letter
    for i, c in enumerate(COLUMNS, 1):
        cell = ws.cell(1, i)
        cell.font, cell.fill, cell.alignment, cell.border = S['hfont'], S['hfill'], S['wrap'], S['border']
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 20)
    for r in rows:
        ws.append([r.get(c, '') for c in COLUMNS])
        for i in range(1, len(COLUMNS) + 1):
            cell = ws.cell(ws.max_row, i)
            cell.font, cell.alignment, cell.border = S['font'], S['wrap'], S['border']
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet('Settings (for scripts)')
    ws.append(['Workflow ID', 'Part', 'Settings JSON (do not edit by hand)'])
    for i in range(1, 4):
        ws.cell(1, i).font, ws.cell(1, i).fill = S['hfont'], S['hfill']
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['C'].width = 120
    for wid, cfg in machine:
        s = json.dumps(cfg, sort_keys=True, separators=(',', ':'))
        for n, k in enumerate(range(0, len(s), CHUNK), 1):
            ws.append([wid, n, s[k:k + CHUNK]])
    wb.save(path)


def read_master(path):
    """{workflow id: {'row': {column: text}, 'cfg': settings dict}} plus the Read Me pairs."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    out = {}
    it = wb['Master Rules'].iter_rows(values_only=True)
    hdr = list(next(it))
    for r in it:
        if r and r[1]:
            d = {h: ('' if v is None else str(v)) for h, v in zip(hdr, r)}
            out[d['Workflow ID']] = {'row': d}
    parts = {}
    it = wb['Settings (for scripts)'].iter_rows(values_only=True)
    next(it)
    for wid, n, s in it:
        if wid:
            parts.setdefault(str(wid), []).append((int(n), s))
    for wid, ps in parts.items():
        out.setdefault(wid, {'row': {}})['cfg'] = json.loads(''.join(s for _, s in sorted(ps)))
    readme = [(a, b) for a, b in wb['Read Me'].iter_rows(min_row=3, values_only=True) if a]
    return out, readme


# ------------------------------------------------------------------ compare
def compare(master_cfg, live_cfg, team_cfg):
    """Differences between the master settings and the live workflow, per compared column.
    Both sides are described with the same labels, so a renamed property label is not reported as a change."""
    a, b = row(master_cfg, team_cfg), row(trim(live_cfg), team_cfg)
    diffs = []
    for col in COMPARED:
        if a[col] == b[col]:
            continue
        if col in MULTILINE:
            old, new = a[col].split('\n'), b[col].split('\n')
            removed = [l for l in old if l not in new]
            added = [l for l in new if l not in old]
            if not removed and not added:
                diffs.append({'column': col, 'was': 'same lines, different order', 'now': '\n'.join(new)})
            else:
                diffs.append({'column': col, 'was': '\n'.join(removed) or '(nothing)', 'now': '\n'.join(added) or '(removed)'})
        else:
            diffs.append({'column': col, 'was': a[col], 'now': b[col]})
    return diffs


# ------------------------------------------------------------------ commands
def _load_work(work):
    if not os.path.exists(os.path.join(work, 'COMPLETE')):
        sys.exit(f'STOP: {work} is incomplete (fetch.py did not finish).')
    describe.init(work)
    return json.load(open(os.path.join(work, 'configs.json')))


def cmd_build(A):
    T = load_team(A.team)
    out = master_path(T)
    if os.path.exists(out) and not A.overwrite:
        sys.exit(f'STOP: {out} already exists. The master rules file is never rebuilt during an audit; '
                 'use "update" for confirmed changes, or --overwrite only if the user asked for a fresh baseline.')
    cfgs = _load_work(A.work)
    purposes = json.load(open(A.purposes)) if A.purposes else {}
    notes = json.load(open(A.notes)) if A.notes else {}
    ids = [w for w in cfgs]
    ids.sort(key=lambda w: (not cfgs[w].get('isEnabled'), cfgs[w]['name'].lower()))
    machine = [(w, trim(cfgs[w])) for w in ids]
    rows = [row(m, T, purposes.get(w, ''), notes.get(w, '')) for w, m in machine]
    snap = max((os.path.getmtime(os.path.join(A.work, 'configs.json')),))
    readme = [
        ('Team', T['team_name']),
        ('Made on', A.made_on or dt.date.today().isoformat()),
        ('Settings as of', dt.datetime.fromtimestamp(snap, dt.timezone.utc).strftime('%Y-%m-%d %H:%M UTC') + ' (saved settings snapshot)'),
        ('Where it came from', A.source_note or 'Built from a saved snapshot of the workflow settings.'),
        ('What it is for', 'The source of truth for "Did anyone change the workflows?". The weekly audit compares every live workflow with its row here.'),
        ('Rules for this file', 'Never rebuild it from the live workflows during an audit; that would hide the changes the audit is meant to catch. '
                                'A row is updated only when WLS confirms the change was planned (master.py update), and the change is logged below.'),
        ('Sheets', 'Master Rules: one row per workflow, readable. Settings (for scripts): the same settings in machine form, used by the scripts to compare '
                   'and to work out which facilities should get each task. Task descriptions are stored as a fingerprint (description id) only.'),
        ('Compared columns', ', '.join(COMPARED)),
        ('Workflows', f'{len(rows)} ({sum(1 for r in rows if r["On/Off"] == "ON")} switched on, {sum(1 for r in rows if r["On/Off"] == "OFF")} switched off)'),
    ]
    for k in T.get('known_issues', []):
        readme.append((f'Known issue {k["id"]}', k['text']))
    for n in T.get('not_readable', []):
        readme.append(('Not in this file', f'{n["name"]} (known issue {n.get("known_issue", "")}): its settings could not be read.'))
    readme.append(('Change log', f'{A.made_on or dt.date.today().isoformat()}: file created.'))
    write_workbook(out, T, rows, machine, readme)
    print(f'Wrote {out}: {len(rows)} workflows')


def cmd_update(A):
    T = load_team(A.team)
    path = master_path(T)
    M, readme = read_master(path)
    cfgs = _load_work(A.work)
    for w in A.ids:
        if w not in cfgs:
            sys.exit(f'STOP: workflow {w} is not in {A.work}.')
    machine, rows = [], []
    for w, e in M.items():
        cfg = trim(cfgs[w]) if w in A.ids else e['cfg']
        machine.append((w, cfg))
        rows.append(row(cfg, T, e['row'].get('Purpose', ''), e['row'].get('Notes', '')))
    for w in A.ids:
        if w not in M:
            machine.append((w, trim(cfgs[w])))
            rows.append(row(trim(cfgs[w]), T, '', 'Added after the first build.'))
    readme = list(readme)
    readme.append(('Change log', f'{dt.date.today().isoformat()}: updated {", ".join(A.ids)}. {A.reason}'))
    write_workbook(path, T, rows, machine, readme)
    print(f'Updated {path}: {", ".join(A.ids)}')


def cmd_remove(A):
    """Remove rows for workflows the user confirmed were deleted on purpose, and log it."""
    T = load_team(A.team)
    path = master_path(T)
    _load_work(A.work)                      # labels for the readable columns (properties, lists, owners)
    M, readme = read_master(path)
    missing = [w for w in A.ids if w not in M]
    if missing:
        sys.exit(f'STOP: not in the master rules file: {", ".join(missing)}')
    names = {w: M[w]['row'].get('Workflow Name') for w in A.ids}
    keep = [(w, e) for w, e in M.items() if w not in A.ids]
    rows = [row(e['cfg'], T, e['row'].get('Purpose', ''), e['row'].get('Notes', '')) for w, e in keep]
    machine = [(w, e['cfg']) for w, e in keep]
    readme = list(readme)
    readme.append(('Change log', f'{dt.date.today().isoformat()}: removed {"; ".join(f"{n} ({w})" for w, n in names.items())}. {A.reason}'))
    write_workbook(path, T, rows, machine, readme)
    print(f'Removed from {path}: {", ".join(A.ids)}')


def cmd_diff(A):
    T = load_team(A.team)
    M, _ = read_master(master_path(T))
    cfgs = _load_work(A.work)
    for w, e in M.items():
        if w not in cfgs:
            print(f'## {e["row"].get("Workflow Name")} ({w}): NOT RETURNED by the API')
            continue
        d = compare(e['cfg'], cfgs[w], T)
        if d:
            print(f'## {e["row"].get("Workflow Name")} ({w})')
            for x in d:
                print(f'- {x["column"]}\n  was: {x["was"]}\n  now: {x["now"]}')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    sp = ap.add_subparsers(dest='cmd', required=True)
    b = sp.add_parser('build')
    b.add_argument('--team', required=True)
    b.add_argument('--work', required=True)
    b.add_argument('--purposes')
    b.add_argument('--notes')
    b.add_argument('--source-note')
    b.add_argument('--made-on')
    b.add_argument('--overwrite', action='store_true')
    u = sp.add_parser('update')
    u.add_argument('--team', required=True)
    u.add_argument('--work', required=True)
    u.add_argument('--ids', nargs='+', required=True)
    u.add_argument('--reason', required=True)
    r = sp.add_parser('remove')
    r.add_argument('--team', required=True)
    r.add_argument('--work', required=True)
    r.add_argument('--ids', nargs='+', required=True)
    r.add_argument('--reason', required=True)
    d = sp.add_parser('diff')
    d.add_argument('--team', required=True)
    d.add_argument('--work', required=True)
    A = ap.parse_args()
    {'build': cmd_build, 'update': cmd_update, 'remove': cmd_remove, 'diff': cmd_diff}[A.cmd](A)
